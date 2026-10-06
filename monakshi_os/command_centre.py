from __future__ import annotations

from datetime import datetime, timezone
from io import BytesIO
import re
from typing import Any, BinaryIO

from openpyxl import load_workbook

from .db import connect, init_db
from .source_canon import upsert_decision, upsert_source
from .state_io import normalize_status


def _version_from_name(name: str) -> str:
    match = re.search(r"(?:^|[_\-\s])v(\d+(?:\.\d+)*)", name or "", re.I)
    return f"v{match.group(1)}" if match else ""


def _as_stream(payload: bytes | bytearray | BinaryIO):
    if isinstance(payload, (bytes, bytearray)):
        return BytesIO(payload)
    return payload


def _clean(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _sheet_rows(ws):
    return [list(row) for row in ws.iter_rows(values_only=True)]


def _upsert_task(row: dict[str, str], db_path: str) -> None:
    source_id = row.get("ID", "")
    if not source_id:
        return
    with connect(db_path) as con:
        con.execute(
            """INSERT INTO launch_tasks(
                source_id, area, task, status, blocker, owner, priority, due_gate,
                dependency, next_action, evidence, notes
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            ON CONFLICT(source_id) DO UPDATE SET
                area=excluded.area,
                task=excluded.task,
                status=excluded.status,
                blocker=excluded.blocker,
                owner=excluded.owner,
                priority=excluded.priority,
                due_gate=excluded.due_gate,
                dependency=excluded.dependency,
                next_action=excluded.next_action,
                evidence=excluded.evidence,
                notes=excluded.notes""",
            (
                source_id,
                row.get("Track", "General") or "General",
                row.get("Task", ""),
                normalize_status(row.get("Status", "")),
                "",
                row.get("Owner", "Founder") or "Founder",
                row.get("Priority", ""),
                row.get("Due / Gate", ""),
                row.get("Dependency", ""),
                row.get("Next Action", ""),
                row.get("Evidence / Output", ""),
                row.get("Notes", ""),
            ),
        )


def _dashboard_decisions(rows: list[list[Any]], source_id: str, effective_date: str) -> list[dict[str, Any]]:
    mapping = {
        "Launch status": ("launch", "launch_status"),
        "Target model": ("operations", "operating_model"),
        "Primary gateway": ("commerce", "primary_gateway"),
        "Backup gateway": ("commerce", "backup_gateway"),
    }
    decisions = []
    for row in rows:
        if not row:
            continue
        label = _clean(row[0])
        if label not in mapping:
            continue
        value = _clean(row[1] if len(row) > 1 else "")
        if not value:
            continue
        domain, key = mapping[label]
        decisions.append(
            {
                "decision_id": f"{source_id}:dashboard:{key}",
                "domain": domain,
                "decision_key": key,
                "decision_value": value,
                "status": "working",
                "source_id": source_id,
                "effective_date": effective_date,
                "rationale": f"Imported from controlling Command Centre Dashboard: {label}.",
            }
        )
    return decisions


def import_command_centre(
    payload: bytes | bytearray | BinaryIO,
    filename: str = "House_of_Monakshi_Launch_Command_Centre.xlsx",
    db_path: str = "data/monakshi.db",
    source_id: str = "command-centre-current",
    source_location: str = "PRIVATE_DRIVE_OR_LOCAL",
) -> dict[str, Any]:
    init_db(db_path)
    workbook = load_workbook(_as_stream(payload), read_only=True, data_only=False)
    now = datetime.now(timezone.utc).isoformat()
    version = _version_from_name(filename)

    upsert_source(
        {
            "source_id": source_id,
            "name": "Current Launch Command Centre",
            "source_type": "spreadsheet",
            "authority_rank": 2,
            "version": version,
            "location": source_location,
            "is_controlling": True,
            "last_verified_at": now,
            "freshness_days": 14,
            "status": "active",
            "notes": f"Imported from {filename}. Workbook contents stay private.",
        },
        db_path,
    )

    task_count = 0
    decisions: list[dict[str, Any]] = []

    if "Master Tasks" in workbook.sheetnames:
        rows = _sheet_rows(workbook["Master Tasks"])
        if rows:
            headers = [_clean(x) for x in rows[0]]
            for values in rows[1:]:
                record = {headers[i]: _clean(values[i] if i < len(values) else "") for i in range(len(headers))}
                if not record.get("ID") or not record.get("Task"):
                    continue
                _upsert_task(record, db_path)
                task_count += 1

                if record["ID"] == "S01":
                    match = re.search(r"Launch Five locked:\s*(.+)$", record.get("Task", ""), re.I)
                    if match:
                        decisions.append(
                            {
                                "decision_id": f"{source_id}:task:S01:launch_suppliers",
                                "domain": "suppliers",
                                "decision_key": "launch_suppliers",
                                "decision_value": match.group(1).strip(),
                                "status": "working",
                                "source_id": source_id,
                                "effective_date": now,
                                "rationale": "Imported from Master Task S01.",
                            }
                        )

                if record["ID"] == "L01":
                    decisions.append(
                        {
                            "decision_id": f"{source_id}:task:L01:legal_structure",
                            "domain": "legal",
                            "decision_key": "legal_structure_working_assumption",
                            "decision_value": "Sole proprietor trading as House of Monakshi",
                            "status": "working",
                            "source_id": source_id,
                            "effective_date": now,
                            "rationale": "Master Task L01 explicitly uses sole proprietorship as the working assumption pending CA confirmation.",
                        }
                    )

                if record["ID"] == "Q05" and record.get("Notes"):
                    shelf = re.search(r"Target soft-launch shelf:\s*([^\.]+)", record["Notes"], re.I)
                    if shelf:
                        decisions.append(
                            {
                                "decision_id": f"{source_id}:task:Q05:soft_launch_shelf",
                                "domain": "products",
                                "decision_key": "soft_launch_shelf",
                                "decision_value": shelf.group(1).strip(),
                                "status": "working",
                                "source_id": source_id,
                                "effective_date": now,
                                "rationale": "Imported from Master Task Q05 notes.",
                            }
                        )

    if "Dashboard" in workbook.sheetnames:
        decisions.extend(_dashboard_decisions(_sheet_rows(workbook["Dashboard"]), source_id, now))

    for decision in decisions:
        upsert_decision(decision, db_path)

    workbook.close()
    return {
        "source_id": source_id,
        "version": version,
        "tasks": task_count,
        "decisions": len(decisions),
        "sheets_seen": workbook.sheetnames if hasattr(workbook, "sheetnames") else [],
    }
